import contextlib
import copy
import io
import json
from collections import Counter
from pathlib import Path
import sys
import unittest
from uuid import uuid4

from openpyxl import load_workbook

TO_EXEL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TO_EXEL))
import ParseJsonToExel as export
from ParseJsonToExel import parse_data
from search import CHASSIS, MANAGER, SYSTEM


@contextlib.contextmanager
def output_folder():
    # Python 3.13 TemporaryDirectory applies an owner-only Windows ACL, which
    # prevents a restricted test process from accessing its own test output.
    folder = Path(__file__).resolve().parent / ("output_" + uuid4().hex)
    folder.mkdir()
    try:
        yield folder
    finally:
        (folder / export.OUTPUT_PATH).unlink(missing_ok=True)
        folder.rmdir()


class DellExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.files = sorted((TO_EXEL / "examples").glob("*.json"))
        if len(cls.files) != 2:
            raise AssertionError("Both Dell example files are required")
        cls.raw = [json.loads(path.read_text(encoding="utf-8")) for path in cls.files]

    def test_inventory_counts_for_both_idrac_schemas(self):
        fields = ("Processors", "Memory", "NetworkAdapters", "SmartStorage", "PhysicalDisks",
                  "PowerSupplies", "Fans", "StorageEnclosures")
        for raw, counts in zip(self.raw, [(2, 8, 3, 3, 2, 0, 0, 1), (2, 16, 2, 3, 2, 0, 0, 1)]):
            data = parse_data(raw)
            self.assertEqual(tuple(len(data[field]) for field in fields), counts)
            self.assertNotIn("TrustedModules", data)  # BIOS explicitly says no TPM.

    def test_values_match_source_responses(self):
        for raw in self.raw:
            data = parse_data(raw)
            self.assertEqual(data["SKU"], raw[SYSTEM]["PartNumber"])
            self.assertEqual(data["SerialNumber"], raw[SYSTEM]["SerialNumber"])
            self.assertEqual(data["iDRACVersion"], raw[MANAGER]["FirmwareVersion"])
            self.assertEqual(data["PowerRegulatorMode"], raw[SYSTEM + "/Bios"]["Attributes"]["SysProfile"])
            self.assertEqual(data["PowerAutoOn"], raw[SYSTEM + "/Bios"]["Attributes"]["AcPwrRcvry"])
            for memory in data["Memory"]:
                source = raw[memory["@odata.id"]]
                self.assertEqual(memory["SizeMB"], source["CapacityMiB"])
                self.assertEqual(memory["Frequency"], source["OperatingSpeedMhz"])
                self.assertEqual(memory["DIMMStatus"], source["Status"]["Health"])
            for disk in data["PhysicalDisks"]:
                source = raw[disk["@odata.id"]]
                self.assertEqual(disk["FirmwareVersion"], source["Revision"])
                self.assertAlmostEqual(disk["CapacityGB"], source["CapacityBytes"] / 1e9)
                self.assertEqual(disk["InterfaceSpeedMbps"], source["NegotiatedSpeedGbs"] * 1000)
                # Both sample exports put a serialized PPID in PartNumber.
                self.assertEqual(disk["PartNumber"], source["Model"])
            for adapter in data["NetworkAdapters"]:
                source = raw[adapter["@odata.id"]]
                self.assertEqual(adapter["FirmwareVersion"], source["Controllers"][0]["FirmwarePackageVersion"])

    def test_controller_identity_from_linked_pcie_device(self):
        data = parse_data(self.raw[0])
        raid = next(c for c in data["SmartStorage"] if c["Model"] == "PERC H330 Mini")
        self.assertEqual(raid["PartNumber"], "07G4YN")
        self.assertEqual(raid["SerialNumber"], "CNFCP007C800CQ")
        self.assertEqual(raid["CacheMemorySizeMiB"], 0)
        self.assertEqual(raid["CurrentOperatingMode"], "RAID")

    def test_backplane_does_not_inherit_controller_identity(self):
        for raw in self.raw:
            data = parse_data(raw)
            for enclosure in data["StorageEnclosures"]:
                source = raw[enclosure["@odata.id"]]
                self.assertEqual(enclosure.get("SerialNumber"), source.get("SerialNumber"))
                self.assertEqual(enclosure.get("PartNumber"), source.get("PartNumber"))
                self.assertEqual(enclosure["ComponentName"], "Дисковая корзина (backplane)")

    def test_enclosure_preserves_own_identity_and_external_category(self):
        from getFromJSON.getArrayControllers import get_storage_enclosures
        path = "/redfish/v1/Chassis/Enclosure.External.1"
        device_path = CHASSIS + "/PCIeDevices/1"
        raw = {path: {"@odata.id": path, "Model": "External enclosure", "SerialNumber": "ENC-SN",
                      "PartNumber": "ENC-PN", "Links": {"PCIeDevices": [{"@odata.id": device_path}]}},
               device_path: {"SerialNumber": "RAID-SN", "PartNumber": "RAID-PN"}}
        enclosure = get_storage_enclosures(raw)[0]
        self.assertEqual(enclosure["SerialNumber"], "ENC-SN")
        self.assertEqual(enclosure["PartNumber"], "ENC-PN")
        self.assertNotIn("ComponentName", enclosure)
        raw[path]["SerialNumber"] = None
        raw[path]["PartNumber"] = None
        enclosure = get_storage_enclosures(raw)[0]
        self.assertIsNone(enclosure["SerialNumber"])
        self.assertIsNone(enclosure["PartNumber"])

    def test_qme_identity_from_own_oem_ports(self):
        data = parse_data(self.raw[0])
        adapters = [n for n in data["NetworkAdapters"] if n["Id"] == "FC.Mezzanine.2B"]
        self.assertEqual(len(adapters), 1)
        self.assertEqual(adapters[0]["PartNumber"], "QME2662")
        self.assertEqual(adapters[0]["Model"], "QME2662")
        self.assertEqual(adapters[0]["Name"], "QME2662")
        self.assertIsNone(adapters[0]["SerialNumber"])
        raw = copy.deepcopy(self.raw[0])
        oem_path = next(k for k in raw if "/Oem/Dell/DellFC/" in k)
        raw[oem_path]["PartNumber"] = "REAL-PN"
        raw[oem_path]["SerialNumber"] = "REAL-SN"
        data = parse_data(raw)
        adapter = next(n for n in data["NetworkAdapters"] if n["Id"] == "FC.Mezzanine.2B")
        self.assertEqual(adapter["PartNumber"], "REAL-PN")
        self.assertEqual(adapter["SerialNumber"], "REAL-SN")
        # Another card's OEM identity and transceiver P/N are not adapter P/N.
        raw[oem_path]["PartNumber"] = None
        raw[oem_path]["TransceiverPartNumber"] = "TRANSCEIVER"
        data = parse_data(raw)
        adapter = next(n for n in data["NetworkAdapters"] if n["Id"] == "FC.Mezzanine.2B")
        self.assertEqual(adapter["PartNumber"], "QME2662")
        self.assertEqual(next(n["PartNumber"] for n in data["NetworkAdapters"]
                              if n["Id"] == "NIC.Integrated.1"), "0XWKGY")

    def test_memory_status_precedence_and_missing_values(self):
        path = SYSTEM + "/Memory/DIMM.Socket.A1"
        raw = {SYSTEM: {"Model": "Dell"}, path: {
            "@odata.id": path, "Name": "DIMM A1", "DIMMStatus": "Explicit",
            "Status": {"Health": "Critical", "State": "Disabled"}}}
        self.assertEqual(parse_data(raw)["Memory"][0]["DIMMStatus"], "Explicit")
        del raw[path]["DIMMStatus"]
        self.assertEqual(parse_data(raw)["Memory"][0]["DIMMStatus"], "Critical")
        raw[path]["Status"] = {"State": "Disabled"}
        self.assertEqual(parse_data(raw)["Memory"][0]["DIMMStatus"], "Disabled")
        raw[path]["Status"] = {}
        self.assertIsNone(parse_data(raw)["Memory"][0]["DIMMStatus"])
        raw[path]["Oem"] = {"Dell": {"DellMemory": {"DIMMStatus": "OEM status"}}}
        self.assertEqual(parse_data(raw)["Memory"][0]["DIMMStatus"], "OEM status")

    def test_disk_model_falls_back_only_without_part_number(self):
        raw = copy.deepcopy(self.raw[0])
        path = next(k for k, v in raw.items() if v.get("@odata.type", "").endswith(".Drive"))
        for missing in (None, ""):
            raw[path]["PartNumber"] = missing
            data = parse_data(raw)
            disk = next(d for d in data["PhysicalDisks"] if d["@odata.id"] == path)
            self.assertEqual(disk["PartNumber"], raw[path]["Model"])
        del raw[path]["PartNumber"]
        data = parse_data(raw)
        disk = next(d for d in data["PhysicalDisks"] if d["@odata.id"] == path)
        self.assertEqual(disk["PartNumber"], raw[path]["Model"])

    def test_disk_ppid_is_rejected_but_real_part_number_is_preserved(self):
        from getFromJSON.getArrayControllers import get_array_controllers
        path = SYSTEM + "/Storage/Drives/Disk.1"
        raw = {path: {"@odata.id": path, "Name": "Disk", "Model": "HUC101860CSS200"}}
        disk = raw[path]
        for part in ("TH-06DWVP-HGT00-85M-4316-A00", "TW0919J9ITT0081302D4A00"):
            disk["PartNumber"] = part
            self.assertEqual(get_array_controllers(raw)[1][0]["PartNumber"], disk["Model"])
        for part in ("06DWVP", "919J9", "HUC101860CSS200", "GENERIC-PART-123"):
            disk["PartNumber"] = part
            self.assertEqual(get_array_controllers(raw)[1][0]["PartNumber"], part)
        disk["PartNumber"] = "serialized-id"
        disk["Oem"] = {"Dell": {"DellPhysicalDisk": {"PPID": "serialized-id", "PartNumber": "06DWVP"}}}
        self.assertEqual(get_array_controllers(raw)[1][0]["PartNumber"], "06DWVP")
        del disk["Oem"]["Dell"]["DellPhysicalDisk"]["PartNumber"]
        self.assertEqual(get_array_controllers(raw)[1][0]["PartNumber"], disk["Model"])
        disk["Model"] = None
        self.assertIsNone(get_array_controllers(raw)[1][0]["PartNumber"])

    def test_legacy_controllers_work_without_modern_resources(self):
        raw = {k: v for k, v in self.raw[0].items() if "/Controllers" not in k}
        data = parse_data(raw)
        self.assertEqual(len(data["SmartStorage"]), 3)
        self.assertIn("07G4YN", [c.get("PartNumber") for c in data["SmartStorage"]])

    def test_repeated_fragments_and_network_views_are_not_counted_twice(self):
        raw = copy.deepcopy(self.raw[0])
        for path, node in list(raw.items()):
            if "#" not in path:
                raw[path + "#/duplicate"] = node
            if path.startswith(CHASSIS + "/NetworkAdapters/"):
                raw[path.replace(CHASSIS, SYSTEM)] = node
        before, after = parse_data(self.raw[0]), parse_data(raw)
        for field in ("Processors", "Memory", "NetworkAdapters", "SmartStorage", "PhysicalDisks"):
            self.assertEqual(before[field], after[field])

    def test_missing_and_unrelated_input(self):
        for raw in ({}, [], {SYSTEM: {"error": {"message": "unavailable"}}}):
            self.assertIsNone(parse_data(raw))
        self.assertIsNone(parse_data({"/redfish/v1/Systems/1": {"Model": "HPE"}}))
        data = parse_data({SYSTEM: {"Model": "Dell"}})
        self.assertIsNone(data["iDRACVersion"])
        self.assertEqual(data["PhysicalDisks"], [])
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertIsNone(export.print_parser(TO_EXEL / "does-not-exist.json"))

    def test_absent_failed_zero_and_false_values(self):
        raw = copy.deepcopy(self.raw[0])
        for path, node in raw.items():
            if "/Processors/CPU.Socket." in path and path.count("/") == 6:
                node["Status"] = {"State": "Absent"}
            if node.get("@odata.type", "").endswith(".Memory"):
                node["Status"] = {"State": "Disabled", "Health": "Critical"}
        raw[SYSTEM + "/Bios"]["Attributes"]["AcPwrRcvry"] = False
        disk_path = next(k for k, v in raw.items() if v.get("@odata.type", "").endswith(".Drive"))
        raw[disk_path]["CapacityBytes"] = 0
        raw[disk_path]["NegotiatedSpeedGbs"] = 0
        data = parse_data(raw)
        self.assertFalse(data["PowerAutoOn"])
        self.assertEqual(len(data["Memory"]), 8)
        self.assertEqual(data["Memory"][0]["Health"], "Critical")
        self.assertEqual(data["Memory"][0]["DIMMStatus"], "Critical")
        self.assertTrue(any(d["CapacityGB"] == 0 and d["InterfaceSpeedMbps"] == 0 for d in data["PhysicalDisks"]))
        from openpyxl import Workbook
        wb = Workbook()
        end, _ = export.write_proc_info(wb.active, data, 1, 1)
        self.assertEqual(end, 1)  # Empty sockets must not become summary CPUs.

    def test_psu_and_fan_legacy_and_modern_paths(self):
        psu = {"Name": "PSU 1", "SerialNumber": "PSU-SN", "Status": {"Health": "Warning"}}
        fan = {"Name": "Fan 1", "Status": {"State": "Enabled"}}
        raw = {SYSTEM: {"Model": "Dell"}, CHASSIS + "/Power": {"PowerSupplies": [psu]},
               CHASSIS + "/Thermal": {"Fans": [fan, {"Name": "Fan 2", "Status": {"State": "Absent"}}]}}
        for modern in (False, True):
            if modern:
                raw[CHASSIS + "/PowerSubsystem/PowerSupplies/1"] = psu
                raw[CHASSIS + "/ThermalSubsystem/Fans/1"] = fan
            data = parse_data(raw)
            self.assertEqual(len(data["PowerSupplies"]), 1)
            self.assertEqual(len(data["Fans"]), 1)
            self.assertEqual(data["PowerSupplies"][0]["Health"], "Warning")

    def test_excel_roundtrip_layout_grouping_and_repeated_export(self):
        expected_headers = ("№", "S/N", "P/N", "Наименование", "Spare P/N", "Description",
                            "Quantity", "Комментарий", "Версия прошивки", "State", "Health",
                            "HealthRollup", "DIMMStatus", "Политика питания", "Включение после потери питания")
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            for _ in range(2):
                export._parseJSONToExel(self.files + self.files, folder)
                wb = load_workbook(Path(folder) / export.OUTPUT_PATH)
                try:
                    self.assertEqual(wb.sheetnames, ["Аудит", "Группировка_PN"])
                    ws = wb["Аудит"]
                    self.assertEqual(next(ws.values), expected_headers)
                    rows = list(ws.values)[1:]
                    servers = [row for row in rows if row[0] is not None]
                    self.assertEqual(len(servers), 2)
                    self.assertTrue(all("iDRAC:" in row[8] for row in servers))
                    components = [row for row in rows if row[6] == 1]
                    self.assertEqual(Counter(row[3] for row in components), {
                        "Процессор": 4, "Память": 24, "Сетевой адаптер": 5,
                        "RAID контроллер": 6, "Диск": 4, "Дисковая корзина (backplane)": 2})
                    raid_rows = [row for row in components if row[1] == "CNFCP007C800CQ"]
                    self.assertEqual(len(raid_rows), 1)
                    self.assertEqual(raid_rows[0][2:4], ("07G4YN", "RAID контроллер"))
                    summary = wb["Группировка_PN"]
                    self.assertEqual([row[:3] for row in summary.values if row[0] == "07G4YN"],
                                     [("07G4YN", "RAID контроллер", 1)])
                    self.assertEqual(sum(row[2] for row in summary.values if isinstance(row[2], int)), len(components))
                    memory_rows = [row for row in components if row[3] == "Память"]
                    self.assertTrue(all(row[12] == "OK" for row in memory_rows))
                    self.assertEqual(len([row for row in components if row[2] == "QME2662"]), 1)
                    expected_disks = {node["SerialNumber"]: node["Model"]
                                      for raw in self.raw for node in raw.values()
                                      if node.get("@odata.type", "").endswith(".Drive")}
                    self.assertEqual({row[1]: row[2] for row in components if row[3] == "Диск"},
                                     expected_disks)
                    for part_number, count in Counter(expected_disks.values()).items():
                        self.assertIn((part_number, "Диск", count), [row[:3] for row in summary.values])
                    for row in ws.iter_rows(min_row=2):
                        if row[3].value == "Память":
                            self.assertEqual(row[12].fill.fgColor.rgb, "00C6EFCE")
                    self.assertEqual(ws["K2"].value, self.raw[0][SYSTEM]["Status"]["Health"])
                    self.assertEqual(ws["K2"].fill.fill_type, "solid")
                    self.assertEqual(ws.column_dimensions["I"].width, 30)
                    self.assertEqual(ws.row_dimensions[1].height, 42)
                finally:
                    wb.close()

    def test_all_project_imports_are_inside_dell(self):
        import main
        self.assertEqual(main._parseJSONToExel, export._parseJSONToExel)
        names = ["main", "ParseJsonToExel", "search", "diagnostic_sheets", "getAllJSONs"]
        names += ["getFromJSON." + path.stem for path in (TO_EXEL / "getFromJSON").glob("*.py")]
        for name in names:
            with self.subTest(module=name):
                path = Path(sys.modules[name].__file__).resolve()
                self.assertTrue(path.is_relative_to(TO_EXEL), str(path))


if __name__ == "__main__":
    unittest.main()
