import contextlib
import copy
import io
import json
from pathlib import Path
import sys
import unittest
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

TO_EXEL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TO_EXEL))
import ParseJsonToExel as export
from diagnostic_sheets import format_status_colors
from fleet_data import firmware_groups, health_issues
from fleet_sheets import write_fleet_sheets
from getFromJSON.getDiagnostics import get_diagnostics
from getFromJSON.getPower import get_power_consumption
from getFromJSON.getSystem import get_system
from search import MANAGER, SYSTEM
from sample_files import reference_files

POWER = "/redfish/v1/Chassis/System.Embedded.1/Power"


@contextlib.contextmanager
def output_folder():
    # Avoid TemporaryDirectory's owner-only Windows ACL under restricted runners.
    folder = Path(__file__).resolve().parent / ("output_" + uuid4().hex)
    folder.mkdir()
    try:
        yield folder
    finally:
        for file in folder.iterdir():
            file.unlink()
        folder.rmdir()


def table_rows(ws, name):
    left, top, right, bottom = range_boundaries(ws.tables[name].ref)
    return list(ws.iter_rows(min_row=top + 1, max_row=bottom, min_col=left, max_col=right, values_only=True))


class FleetExportTests(unittest.TestCase):
    def test_power_sources_zero_missing_and_overlapping_domains(self):
        control = {"MemberId": "0", "PowerConsumedWatts": 0, "PowerCapacityWatts": 800}
        for body in ({"PowerControl": [control]}, [control], control):
            with self.subTest(body=body):
                self.assertEqual(get_power_consumption({POWER + "#PowerControl": body}), 0)
        raw = {POWER: {"PowerControl": [control]},
               POWER + "#PowerControl/0": {"PowerControl": [{"PowerConsumedWatts": 999}]}}
        self.assertEqual(get_power_consumption(raw), 0)
        raw[POWER]["PowerControl"].append({"MemberId": "1", "PowerConsumedWatts": 50})
        self.assertEqual(get_power_consumption(raw), 0)
        self.assertIsNone(get_power_consumption({POWER: {"PowerSupplies": [{"PowerCapacityWatts": 800}]}}))
        self.assertIsNone(get_power_consumption({POWER: {"PowerControl": [{"PowerConsumedWatts": None}]}}))
        self.assertIsNone(get_power_consumption({}))

    def test_hundred_servers_80_20_shares_and_independent_versions(self):
        servers = [{"SerialNumber": f"SN{i:03}", "Model": "BL460c", "BiosVersion": "A" if i < 80 else "B",
                    "iDRACVersion": "3.19", "ServerHealth": "OK"} for i in range(100)]
        servers += [{"SerialNumber": "OTHER", "Model": "DL380", "BiosVersion": "A", "iDRACVersion": "3.18",
                     "ServerHealth": "Critical"}]
        wb = Workbook()
        write_fleet_sheets(wb, servers)
        ws = wb["Сводка_серверов"]
        self.assertEqual([ws.cell(2, c).value for c in range(1, 6)], [101, 2, 1, 2, 2])
        rows = table_rows(ws, "ServerSummary")
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0][3:5], (80, 0.8))
        self.assertEqual(rows[1][3:5], (20, 0.2))
        self.assertEqual(len(rows[0][5].split(", ")), 80)
        self.assertEqual(ws["E7"].number_format, "0.0%")
        ilo = table_rows(ws, "SummaryiDRAC")
        self.assertEqual(ilo[0][2:5], (100, 100, 1))
        self.assertEqual(wb.active.title, "Сводка_серверов")
        servers[0]["iDRACVersion"] = "3.18"
        wb = Workbook()
        write_fleet_sheets(wb, servers)
        ws = wb["Сводка_серверов"]
        self.assertEqual(len(table_rows(ws, "ServerSummary")), 4)
        self.assertEqual(table_rows(ws, "SummaryBIOS")[0][2], 80)

    def test_three_servers_six_ssds_aliases_and_missing_serials(self):
        servers = []
        for index in range(3):
            disks = [{"Model": "VK000240GWSRQ", "PartNumber": "SSD-PN", "FirmwareVersion": "HPG0",
                      "MediaType": "SSD", "SerialNumber": f"{index}-{slot}" if index < 2 else "",
                      "SourcePath": f"/redfish/v1/Systems/System.Embedded.1/Storage/1/Drives/{slot}"} for slot in range(2)]
            inventory = [dict(disk, ComponentPath=disk["SourcePath"], Field="FirmwareVersion", Version="HPG0") for disk in disks]
            servers.append({"Model": "BL460c", "SerialNumber": f"SN{index}", "PhysicalDisks": disks,
                            "FirmwareInventory": inventory + copy.deepcopy(inventory)})
        groups = firmware_groups(servers)
        self.assertEqual(len(groups), 1)
        key, group = next(iter(groups.items()))
        self.assertEqual(key[:3], ("SSD", "VK000240GWSRQ / SSD-PN", "HPG0"))
        self.assertEqual(len(group["servers"]), 3)
        self.assertEqual(len(group["components"]), 6)
        servers[0]["PhysicalDisks"][0]["FirmwareVersion"] = "HPG1"
        groups = firmware_groups(servers)
        self.assertEqual(sorted(len(g["components"]) for g in groups.values()), [1, 5])

    def test_firmware_microcodes_backups_and_fragment_duplicates(self):
        path = "/redfish/v1/Systems/System.Embedded.1/Processors/1"
        cpu = {"@odata.id": path, "Name": "CPU", "Model": "Xeon", "ProcessorId": {"MicrocodeInfo": "0x12"},
               "Oem": {"Dell": {"MicrocodePatches": [{"CpuId": "CPU-ID", "PatchId": "0x34"}]}}}
        raw = {path: cpu, path + "#Oem": copy.deepcopy(cpu),
               "/redfish/v1/Systems/System.Embedded.1/Processors/2": dict(cpu, **{"@odata.id": "/redfish/v1/Systems/System.Embedded.1/Processors/2"}),
               "/redfish/v1/Systems/System.Embedded.1/Storage/1": {"Name": "RAID", "Model": "P408", "FirmwareVersion": {
                   "Current": {"VersionString": "5.0"}, "Backup": {"VersionString": "4.0"}}}}
        groups = firmware_groups([dict(get_diagnostics(raw), Model="DL380", SerialNumber="SN")])
        versions = {(key[2], key[3]): len(group["components"]) for key, group in groups.items()}
        self.assertEqual(versions, {("0x12", "MicrocodeInfo"): 2,
                                   ("0x34", "Доступный патч CPU, CpuId=CPU-ID"): 2,
                                   ("5.0", "Текущая"): 1, ("4.0", "Backup"): 1})

    def test_health_server_only_and_psu_identification(self):
        raw = {POWER: {"PowerSupplies": [{"Name": "PSU 1", "Model": "800W", "PartNumber": "754381-001",
                                        "SerialNumber": "PSU-SN", "FirmwareVersion": "1.00",
                                        "Status": {"State": "UnavailableOffline", "Health": "Warning"}}]}}
        servers = [{"SerialNumber": "SERVER-ONLY", "Model": "DL380", "ServerHealth": "Critical", "BiosVersion": "BIOS-A"},
                   dict(get_diagnostics(raw), SerialNumber="PSU-SERVER", Model="DL380", ServerHealth="OK")]
        wb = Workbook()
        write_fleet_sheets(wb, servers)
        format_status_colors(wb)
        rows = table_rows(wb["Health"], "HealthIssues")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][2], "Сервер")
        self.assertEqual(rows[0][6], "Critical")
        self.assertEqual(rows[1][2:8], ("Блок питания", "754381-001", "PSU-SN", "UnavailableOffline", "Warning", "1.00"))
        self.assertEqual(wb["Health"]["G2"].fill.fgColor.rgb, "00FFC7CE")
        self.assertEqual(wb["Health"]["G3"].fill.fgColor.rgb, "00FFEB9C")
        self.assertEqual(wb["Сводка_серверов"]["C2"].value, 2)
        self.assertTrue(wb["Health"].conditional_formatting)

    def test_health_nested_sensors_rollup_missing_and_history(self):
        path = "/redfish/v1/Chassis/System.Embedded.1/Thermal"
        thermal = {"@odata.id": path, "Fans": [
            {"Name": "Fan1", "Status": {"Health": "Warning"}},
            {"Name": "Fan2", "Status": {"Health": "OK"}},
            {"Name": "Fan3", "Status": {"Health": None}},
            {"Name": "Fan4", "Status": {"State": "Absent"}},
            {"Name": "Fan5", "Status": {"Health": "Unknown"}}]}
        raw = {path: thermal, path + "#Fans/0": copy.deepcopy(thermal),
               "/redfish/v1/Systems/System.Embedded.1": {"Status": {"Health": "OK", "HealthRollup": "Critical"}},
               "/redfish/v1/Systems/System.Embedded.1/LogServices/IML/Entries/1": {"Status": {"Health": "Critical"}}}
        server = dict(get_diagnostics(raw), SerialNumber="SN", Model="DL380", ServerHealth="OK", ServerHealthRollup="Critical")
        issues = health_issues(server)
        self.assertEqual(len(issues), 3)
        self.assertEqual({item["Health"] for item in issues}, {"Warning", "Unknown", "OK"})
        self.assertTrue(any(item["SourcePath"] == path + "#/Fans/0" for item in issues))

    def test_all_examples_saved_five_sheets_pn_power_and_repeat_run(self):
        files = reference_files()
        self.assertEqual(len(files), 10)
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            export._parseJSONToExel(files + files[:1], folder)
            wb = load_workbook(folder / export.OUTPUT_PATH)
            self.assertEqual(wb.sheetnames, ["Аудит", "Группировка_PN", "Сводка_серверов", "Микрокоды", "Health"])
            self.assertEqual(wb.active.title, "Сводка_серверов")
            audit = wb["Аудит"]
            headers = {cell.value: cell.column for cell in audit[1]}
            self.assertEqual(headers["Dell Service Tag"], 2)
            self.assertEqual(audit.max_column, 17)
            self.assertNotIn("Quantity", headers)
            self.assertEqual(headers["Версия прошивки"], 8)
            self.assertNotIn("Включение после потери питания", headers)
            self.assertEqual([cell.value for cell in audit[1]][13:],
                             ["Redundancy Policy", "Hot Spare", "Primary PSU", "Потребление, Вт"])
            expected = {}
            for file in files:
                data = export.print_parser(file)
                expected.setdefault(data["ServiceTag"], data)
                raw = json.loads(file.read_text(encoding="utf-8"))
                self.assertEqual(data["PowerConsumedWatts"], raw[POWER]["PowerControl"][0].get("PowerConsumedWatts"))
            self.assertEqual(wb["Сводка_серверов"]["A2"].value, len(expected))
            self.assertEqual(sum(row[3] for row in table_rows(wb["Сводка_серверов"], "ServerSummary")), len(expected))
            component_count = 0
            for row in audit.iter_rows(min_row=2):
                if row[0].value is not None:
                    data = expected[row[1].value]
                    self.assertEqual(row[headers["Потребление, Вт"] - 1].value, data["PowerConsumedWatts"])
                    for label, field in (("Redundancy Policy", "RedundancyPolicy"),
                                         ("Hot Spare", "HotSpare"), ("Primary PSU", "PrimaryPSU")):
                        self.assertEqual(row[headers[label] - 1].value, data[field])
                    self.assertTrue(row[1].font.bold)
                    self.assertIsNone(row[1].fill.patternType)
                elif row[2].value is not None:
                    self.assertIsNone(row[1].value)
                    component_count += 1
            pn = wb["Группировка_PN"]
            self.assertEqual([c.value for c in pn[1]], ["P/N", "Компонент", "Количество", "Описание (пример)"])
            self.assertEqual(sum(row[2] for row in pn.iter_rows(values_only=True) if isinstance(row[2], int)), component_count)
            firmware = table_rows(wb["Микрокоды"], "ComponentFirmware")
            self.assertTrue(any(row[0] == "FC HBA" for row in firmware))
            self.assertTrue(any(row[5] == "Previous" for row in firmware))
            self.assertTrue(any(row[5] == "MicrocodeInfo" for row in firmware))
            self.assertTrue(any("BOSS" in (row[1] or "") for row in firmware))
            for ws in (audit, wb["Health"]):
                labels = {cell.value: cell.column for cell in ws[1]}
                for row in ws.iter_rows(min_row=2):
                    for label in ("Health", "HealthRollup"):
                        cell = row[labels[label] - 1]
                        if cell.value in {"Critical", "Warning"}:
                            self.assertEqual(cell.fill.fgColor.rgb, "00FFC7CE" if cell.value == "Critical" else "00FFEB9C")
            wb.close()
            export._parseJSONToExel(files[:1], folder)
            wb = load_workbook(folder / export.OUTPUT_PATH)
            self.assertEqual(wb["Сводка_серверов"]["A2"].value, 1)
            wb.close()

    def test_empty_export_retains_headers_and_zero_metrics(self):
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            export._parseJSONToExel([], folder)
            wb = load_workbook(folder / export.OUTPUT_PATH)
            self.assertEqual(wb["Health"].max_row, 1)
            self.assertEqual(wb["Микрокоды"].max_row, 1)
            self.assertEqual([wb["Сводка_серверов"].cell(2, c).value for c in range(1, 6)], [0] * 5)
            wb.close()

    def test_firmware_groups_across_server_models_without_serial_lists(self):
        servers = []
        for index, model in enumerate(("R650", "R660", "R660")):
            product = f"Broadcom BCM5720 - 00:11:22:33:44:{index:02X}"
            servers.append({"Model": model, "SerialNumber": f"SERVER-{index}", "NetworkAdapters": [
                {"Model": product, "PartNumber": product, "SerialNumber": f"NIC-{index}",
                 "FirmwareVersion": "1.0" if index < 2 else "2.0"}]})
        wb = Workbook()
        write_fleet_sheets(wb, servers)
        ws = wb["Микрокоды"]
        self.assertEqual([cell.value for cell in ws[1]], ["Тип компонента", "Модель / P/N", "Firmware",
                                                         "Компонентов", "Серверов", "Назначение прошивки"])
        self.assertEqual(table_rows(ws, "ComponentFirmware"), [
            ("Сетевой адаптер", "Broadcom BCM5720", "1.0", 2, 2, "Текущая"),
            ("Сетевой адаптер", "Broadcom BCM5720", "2.0", 1, 1, "Текущая")])
        for row in ws.values:
            self.assertFalse(any("SERVER-" in str(value) or "R650" in str(value) or "R660" in str(value)
                                 for value in row))

    def test_license_critical_does_not_flag_healthy_server(self):
        file = next((TO_EXEL / "examples").glob("*1GLZQ34*/redfish_full.json"))
        raw = json.loads(file.read_text(encoding="utf-8"))
        data = export.parse_data(raw)
        self.assertEqual(data["ServerHealth"], "OK")
        self.assertEqual(data["ServerHealthRollup"], "OK")
        licenses = [item for item in data["ComponentStatus"] if "/LicenseService/Licenses/" in item["SourcePath"]]
        self.assertEqual(sum(item["Health"] == "Critical" for item in licenses), 2)
        self.assertEqual(health_issues(data), [])
        wb = Workbook()
        write_fleet_sheets(wb, [data])
        self.assertEqual(wb["Health"].max_row, 1)
        self.assertEqual(wb["Сводка_серверов"]["C2"].value, 0)
        # Missing hardware identifiers must not suppress a real fault.
        data["ComponentStatus"].append({"SourcePath": POWER + "#/PowerSupplies/0",
                                         "Name": "PSU", "Field": "Status", "Health": "Critical"})
        issues = health_issues(data)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["ComponentType"], "Блок питания")

    def test_package_versions_share_one_component_name(self):
        servers = []
        for version in ("1.0", "2.0"):
            path = "/redfish/v1/UpdateService/FirmwareInventory/Installed-1-" + version + "__DriverPack.Embedded.1"
            raw = {path: {"Id": path.rsplit("/", 1)[1], "Name": "Dell OS Driver Pack, " + version,
                          "Version": version}}
            servers.append(get_diagnostics(raw))
        groups = firmware_groups(servers)
        self.assertEqual({key[1] for key in groups}, {"Dell OS Driver Pack"})
        self.assertEqual({key[2] for key in groups}, {"1.0", "2.0"})

    def test_power_policy_uses_current_manager_attributes_and_preserves_missing_values(self):
        path = MANAGER + "/Oem/Dell/DellAttributes/System.Embedded.1"
        raw = {SYSTEM + "/Bios": {"Attributes": {"ServerPwr.1.PSRedPolicy": "unrelated"}},
               path + "/Settings": {"Attributes": {"ServerPwr.1.PSRedPolicy": "pending"}}}
        for attrs, expected in (({}, (None, None, None)),
                                ({"ServerPwr.1.PSRedPolicy": "Not Redundant",
                                  "ServerPwr.1.PSRapidOn": "Disabled",
                                  "ServerPwr.1.RapidOnPrimaryPSU": "PSU2"},
                                 ("Not Redundant", "Disabled", "PSU2")),
                                ({"ServerPwr.1.PSRedPolicy": None, "ServerPwr.1.PSRapidOn": False,
                                  "ServerPwr.1.RapidOnPrimaryPSU": 0}, (None, False, 0))):
            raw[path] = {"Attributes": attrs}
            data = get_system(raw)
            self.assertEqual(tuple(data[key] for key in ("RedundancyPolicy", "HotSpare", "PrimaryPSU")), expected)
            ws = Workbook().active
            export.write_desc_info(ws, 1, 1)
            export.write_server_info(ws, data, 2, 1)
            self.assertEqual(tuple(ws.cell(2, column).value for column in (14, 15, 16)), expected)


if __name__ == "__main__":
    unittest.main()
