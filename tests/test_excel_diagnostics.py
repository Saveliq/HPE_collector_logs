import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "ToExel"))
from ParseJsonToExel import _parseJSONToExel, print_parser, OUTPUT_PATH
from getFromJSON.getDiagnostics import get_diagnostics, component_fields
from getFromJSON.getArrayControllers import get_array_controllers
from getFromJSON.getPower import get_power
from getFromJSON.getNetworkAdapters import get_network_adapters
from diagnostic_sheets import display_value, format_status_colors
from search import stripJSON


class ExampleDiagnosticsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.files = sorted((ROOT / "ToExel" / "examples").glob("*.json"))
        if len(cls.files) != 7:
            raise AssertionError("Expected all seven supplied JSON examples")
        cls.raw = [json.loads(p.read_text(encoding="utf-8")) for p in cls.files]
        with contextlib.redirect_stdout(io.StringIO()):
            cls.parsed = [print_parser(p) for p in cls.files]

    def test_all_examples_and_actual_power_modes(self):
        for data, expected in zip(self.parsed, ["Dynamic", "Dynamic", "Max", "Max", "Dynamic", "Max", "Dynamic"]):
            with self.subTest(serial=data["SerialNumber"]):
                self.assertEqual(data["PowerRegulatorMode"], expected)
                self.assertTrue(data["FirmwareInventory"])
                self.assertTrue(data["Temperatures"])
                self.assertTrue(data["ComponentStatus"])
        self.assertEqual(self.parsed[5]["PowerAutoOn"], "RemainOff")
        self.assertEqual(self.parsed[2]["ServerHealth"], "Critical")
        self.assertEqual(self.parsed[0]["ServerHealth"], "Warning")

    def test_all_thermal_sensors_match_sources_without_fragment_duplicates(self):
        for raw, data in zip(self.raw, self.parsed):
            sensors = raw["/redfish/v1/Chassis/1/Thermal"]["Temperatures"]
            actual = [r for r in data["Temperatures"] if r["Field"] == "ReadingCelsius"]
            self.assertEqual(len(actual), len(sensors))
            for sensor, row in zip(sensors, actual):
                self.assertEqual(row["Name"], sensor["Name"])
                self.assertEqual(row["ReadingCelsius"], sensor["ReadingCelsius"])
                self.assertEqual(row["Health"], sensor.get("Status", {}).get("Health"))
                self.assertEqual(row["UpperThresholdCritical"], sensor.get("UpperThresholdCritical"))

    def test_firmware_covers_oem_components_bios_ilo_and_devices(self):
        for raw, data in zip(self.raw, self.parsed):
            rows = data["FirmwareInventory"]
            versions = [r["Version"] for r in rows]
            self.assertIn(raw["/redfish/v1/Systems/1"]["BiosVersion"], versions)
            self.assertIn(raw["/redfish/v1/Managers/1"]["FirmwareVersion"], versions)
            chassis_fw = raw["/redfish/v1/Chassis/1"]["Oem"]["Hpe"]["Firmware"]
            for name, value in chassis_fw.items():
                matches = [r for r in rows if r["Field"] == "Firmware/" + name + "/Current"]
                self.assertEqual(len(matches), 1)
                self.assertEqual(matches[0]["Version"], value["Current"]["VersionString"])
            for path, device in raw.items():
                if "/Devices/" in path and "FirmwareVersion" in device:
                    self.assertTrue(any(r["SourcePath"] == path and r["Version"] ==
                                        device["FirmwareVersion"]["Current"]["VersionString"] for r in rows))

    def test_every_current_component_status_is_retained(self):
        # Independent traversal: check the source's actual fields, not parser counts.
        for raw, parsed in zip(self.raw, self.parsed):
            def check(node, path):
                if isinstance(node, list):
                    for i, value in enumerate(node):
                        check(value, path + "/" + str(i))
                elif isinstance(node, dict):
                    if isinstance(node.get("Status"), dict):
                        matches = [r for r in parsed["ComponentStatus"] if r["SourcePath"] == path and r["Field"] == "Status"]
                        self.assertEqual(len(matches), 1, path)
                        for field in ("State", "Health", "HealthRollup"):
                            self.assertEqual(matches[0][field], node["Status"].get(field), path)
                    for key, value in node.items():
                        if key not in ("Links", "links", "Actions", "Members", "RelatedItem", "Status") and not key.startswith("@"):
                            check(value, path + "/" + key)
            for source, node in raw.items():
                if "#" not in source and not any(p.lower() in {"settings", "entries"} for p in source.split("/")):
                    # Nested fields use #/ paths; top-level statuses use the resource path.
                    if "Status" in node:
                        matches = [r for r in parsed["ComponentStatus"] if r["SourcePath"] == source and r["Field"] == "Status"]
                        self.assertEqual(len(matches), 1, source)
                    for key, value in node.items():
                        if key not in ("Links", "links", "Actions", "Members", "RelatedItem", "Status") and not key.startswith("@"):
                            check(value, source + "#/" + key)

    def test_disk_temperatures_and_all_fans(self):
        for raw, data in zip(self.raw, self.parsed):
            fan_rows = [r for r in data["ComponentStatus"] if "#/Fans/" in r["SourcePath"] and r["Field"] == "Status"]
            self.assertEqual(len(fan_rows), len(raw["/redfish/v1/Chassis/1/Thermal"]["Fans"]))
            for disk in data["PhysicalDisks"]:
                self.assertTrue(disk["FirmwareVersion"])
                self.assertIsInstance(disk["TemperatureCelsius"], (int, float))

    def test_workbook_roundtrip_and_repeated_runs(self):
        with tempfile.TemporaryDirectory() as directory:
            with contextlib.redirect_stdout(io.StringIO()):
                _parseJSONToExel(self.files + [self.files[0]], directory)
            wb = load_workbook(Path(directory) / OUTPUT_PATH)
            try:
                self.assertEqual(wb.sheetnames, ["Аудит", "Группировка_PN"])
                self.assertEqual(wb["Аудит"].freeze_panes, "D2")
                audit = list(wb["Аудит"].values)
                servers = [r for r in audit[1:] if isinstance(r[0], int)]
                self.assertEqual(len(servers), 7)
                self.assertEqual([r[14] for r in servers], [d["PowerRegulatorMode"] for d in self.parsed])
                self.assertEqual(servers[2][10], "Critical")
                self.assertTrue(any(isinstance(r[12], (int, float)) for r in audit if r[3] == "Диск"))
                self.assertTrue(any(r[8] == "1.00" for r in audit if r[3] == "Блок питания"))
                self.assertFalse(any(c.data_type == "f" for ws in wb for row in ws for c in row))
                self.assertFalse(any(c.value == "Нет данных" for ws in wb for row in ws for c in row))
                self.assertTrue(any(r[8] is None for r in audit if r[3] == "Память"))
                for row in wb["Аудит"].iter_rows(min_row=2):
                    if row[10].value in ("Critical", "Warning"):
                        self.assertEqual(row[10].fill.fgColor.rgb, "00FFC7CE")
                    elif row[10].value == "OK":
                        self.assertEqual(row[10].fill.fgColor.rgb, "00C6EFCE")
                self.assertGreater(wb["Группировка_PN"].max_row, 1)
                # Preserve the earlier layout; only status cells receive colors.
                ws = wb["Аудит"]
                self.assertEqual(ws.row_dimensions[1].height, 42)
                self.assertEqual(ws.column_dimensions["I"].width, 30)
                self.assertNotIn("F", ws.column_dimensions)
                self.assertIsNone(ws["A1"].fill.patternType)
                self.assertIsNone(ws["A2"].fill.patternType)
                memory = next(row for row in ws.iter_rows(min_row=2) if row[3].value == "Память")
                self.assertTrue(memory[5].value.startswith("{'SizeMB':"))
                self.assertNotIn("\n", memory[5].value)
                self.assertFalse(memory[5].alignment.wrap_text)
                self.assertIsNone(memory[5].border.bottom.style)
                self.assertIsNone(ws.row_dimensions[memory[0].row].height)
                self.assertFalse(wb["Группировка_PN"].column_dimensions)
            finally:
                wb.close()
            with contextlib.redirect_stdout(io.StringIO()):
                _parseJSONToExel([self.files[0]], directory)
            wb = load_workbook(Path(directory) / OUTPUT_PATH)
            self.assertEqual(wb.sheetnames, ["Аудит", "Группировка_PN"])
            self.assertEqual(sum(isinstance(r[0], int) for r in list(wb["Аудит"].values)[1:]), 1)
            wb.close()


class DiagnosticEdgeCaseTests(unittest.TestCase):
    def test_missing_values_are_blank_but_zero_and_false_survive(self):
        for value in (None, "", "None", "null", "N/A", "Unknown", " Нет данных ", {}, []):
            self.assertIsNone(display_value(value))
        self.assertEqual(display_value(0), 0)
        self.assertIs(display_value(False), False)
        self.assertEqual(display_value("1.00"), "1.00")

    def test_status_colors_survive_excel_roundtrip(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Аудит"
        ws.append(["Наименование", "Health", "State", "DIMMStatus", "Версия прошивки"])
        ws.append(["Critical", "OK", "Enabled", "GoodInUse", "OK"])
        ws.append(["OK", "Critical", "Unknown", "MapOutError", "Critical"])
        ws.append([None, "Warning", "Absent", "MapOutConfiguration", None])
        ws.append([None, None, None, None, None])
        format_status_colors(wb)
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        saved = load_workbook(stream)
        try:
            ws = saved["Аудит"]
            for ref in ("B2", "C2", "D2"):
                self.assertEqual(ws[ref].fill.fgColor.rgb, "00C6EFCE")
            for ref in ("B3", "D3", "B4", "D4"):
                self.assertEqual(ws[ref].fill.fgColor.rgb, "00FFC7CE")
            for ref in ("A2", "A3", "E2", "E3", "C3", "C4", "B5"):
                self.assertIsNone(ws[ref].fill.patternType)
        finally:
            saved.close()

    def test_network_controller_firmware_and_status_from_second_view(self):
        raw = {
            "/redfish/v1/Chassis/1/NetworkAdapters/AB12": {
                "Name": "Adapter", "SerialNumber": "nic1",
                "Controllers": [{"FirmwarePackageVersion": "07.19.21.00"}]},
            "/redfish/v1/Chassis/1/Devices/1": {
                "Name": "Network adapter", "SerialNumber": "nic1",
                "FirmwareVersion": {"Current": {"VersionString": "7.19.21"}},
                "Status": {"Health": "Warning", "State": "Enabled"}},
        }
        adapters = get_network_adapters(raw)
        self.assertEqual(len(adapters), 1)
        self.assertEqual(adapters[0]["FirmwareVersion"], "07.19.21.00")
        self.assertEqual(adapters[0]["Health"], "Warning")

    def test_power_supply_uses_own_status_and_memory_keeps_oem_status(self):
        fields = component_fields({"Oem": {"Hpe": {"PowerSupplyStatus": {"State": "Ok"}}},
                                   "Status": {"State": "Disabled", "Health": "Critical"}})
        self.assertEqual(fields["State"], "Disabled")
        self.assertEqual(fields["Health"], "Critical")
        self.assertEqual(component_fields({"Oem": {"Hpe": {"DIMMStatus": "MapOutError"}}})["DIMMStatus"], "MapOutError")
        self.assertEqual(component_fields({"DIMMStatus": "NotPresent"})["DIMMStatus"], "NotPresent")

    def test_null_zero_and_false_are_not_confused(self):
        raw = {"/redfish/v1/Chassis/1/Thermal": {"Temperatures": [
            {"Name": "zero", "ReadingCelsius": 0, "Status": {"State": "Enabled"}},
            {"Name": "absent", "ReadingCelsius": None, "Status": {"State": "Absent"}}]},
            "/redfish/v1/Chassis/1/Power": {"Cap": 0, "BrownoutRecoveryEnabled": False}}
        data = get_diagnostics(raw)
        self.assertEqual([r["ReadingCelsius"] for r in data["Temperatures"]], [0, None])
        self.assertEqual({r["Field"]: r["Value"] for r in data["PowerPolicy"]}, {"Cap": 0, "BrownoutRecoveryEnabled": False})
        self.assertEqual(stripJSON({"n": None, "zero": 0, "flag": False, "s": " a "}),
                         {"n": None, "zero": 0, "flag": False, "s": "a"})

    def test_firmware_variants_and_missing_child_health(self):
        raw = {"/redfish/v1/Systems/1": {"FirmwareVersion": {"Current": {"VersionString": "1.00"},
                "Backup": {"VersionString": "0.90"}}, "ControllerBoard": {"Status": {"Health": "OK"}}}}
        data = get_diagnostics(raw)
        self.assertEqual({r["Version"] for r in data["FirmwareInventory"]}, {"1.00", "0.90"})
        fields = component_fields(raw["/redfish/v1/Systems/1"])
        self.assertEqual(fields["FirmwareVersion"], "1.00")
        self.assertIsNone(fields["Health"])

    def test_fragment_only_response_fallback_and_priority(self):
        base = "/redfish/v1/Chassis/1/Thermal"
        body = {"@odata.id": base, "Temperatures": [{"Name": "CPU", "ReadingCelsius": 40}]}
        raw = {base + "#Fans/0": body, base + "#Fans/1": body}
        self.assertEqual(len(get_diagnostics(raw)["Temperatures"]), 1)
        raw[base] = {"Temperatures": [{"Name": "CPU", "ReadingCelsius": 45}]}
        self.assertEqual(get_diagnostics(raw)["Temperatures"][0]["ReadingCelsius"], 45)

    def test_empty_data_and_pending_settings(self):
        self.assertTrue(all(not rows for rows in get_diagnostics({}).values()))
        data = get_diagnostics({"/redfish/v1/Systems/1/Bios/Settings": {"PowerProfile": "Pending"}})
        self.assertFalse(data["PowerPolicy"])
        self.assertEqual(get_power({"/redfish/v1/Chassis/1/Power": {"PowerControl": []}}), [])

    def test_multidigit_raid_controller_id(self):
        base = "/redfish/v1/Systems/1/SmartStorage/ArrayControllers/12"
        controllers, disks = get_array_controllers({base: {"Name": "RAID"}, base + "/DiskDrives/0": {
            "SerialNumber": "disk12", "FirmwareVersion": {"Current": {"VersionString": "HPG0"}},
            "CurrentTemperatureCelsius": 0, "Status": {"Health": "Warning"}}})
        self.assertEqual(len(disks), 1)
        self.assertEqual(disks[0]["SerialNumber"], "disk12")
        self.assertEqual(disks[0]["TemperatureCelsius"], 0)
        self.assertEqual(disks[0]["Health"], "Warning")

    def test_minimal_and_invalid_json(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "server.json"
            path.write_text(json.dumps({"/redfish/v1/Chassis/1": {"Model": "ProLiant", "SerialNumber": "123"}}))
            with contextlib.redirect_stdout(io.StringIO()):
                data = print_parser(path)
            self.assertIsNotNone(data)
            self.assertEqual(data["PowerSupplies"], [])
            for text in ("[]", "null", "broken JSON"):
                path.write_text(text)
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertIsNone(print_parser(path))


if __name__ == "__main__":
    unittest.main()
