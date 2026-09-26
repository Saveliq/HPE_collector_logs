import contextlib
import copy
import io
import json
from collections import Counter
from pathlib import Path
import sys
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "ToExel"))
import ParseJsonToExel as export
from getFromJSON.getArrayControllers import get_array_controllers
from getFromJSON.getNetworkAdapters import get_network_adapters
from getFromJSON.getChassis import get_chassis


WRITERS = (export.write_proc_info, export.write_NIC_info, export.write_RAID_info,
           export.write_MEM_info, export.write_DISK_info, export.write_PSU_info,
           export.write_other_info)


def inventory(data):
    wb = Workbook()
    row = 1
    for writer in WRITERS:
        row, _ = writer(wb.active, data, row, 1)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    saved = load_workbook(stream)
    try:
        return list(saved.active.values) if row > 1 else []
    finally:
        saved.close()


class ComponentInventoryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.files = sorted((ROOT / "ToExel/examples").glob("*.json"))
        cls.raw = [json.loads(path.read_text(encoding="utf-8")) for path in cls.files]
        with contextlib.redirect_stdout(io.StringIO()):
            cls.parsed = [export.print_parser(path) for path in cls.files]

    def test_all_seven_inventory_counts_after_excel_roundtrip(self):
        # Independently audited against installed components in the supplied JSONs.
        categories = ("Процессор", "Сетевой адаптер", "RAID контроллер", "RAID батарея",
                      "Память", "Диск", "Блок питания", "Дисковая полка", "Вентилятор", "USB устройство")
        expected = [(2, 2, 1, 1, 2, 0, 2, 0, 6, 0), (1, 1, 1, 0, 1, 2, 0, 0, 0, 0),
                    (2, 0, 2, 1, 8, 2, 0, 1, 0, 0), (2, 2, 2, 1, 8, 2, 0, 1, 0, 0),
                    (2, 2, 1, 1, 8, 2, 0, 0, 0, 0), (2, 2, 1, 1, 8, 2, 0, 0, 0, 1),
                    (2, 2, 1, 1, 6, 2, 0, 0, 0, 0)]
        for file, data, counts in zip(self.files, self.parsed, expected):
            with self.subTest(file=file.name):
                rows = inventory(data)
                self.assertEqual(Counter(row[3] for row in rows),
                                 Counter({key: count for key, count in zip(categories, counts) if count}))
                self.assertFalse(any(row[9] == "Absent" or row[12] == "NotPresent" for row in rows))

    def test_nvme_boot_controller_and_disks_are_not_duplicated(self):
        data = self.parsed[1]
        self.assertEqual(data["Model"], "Synergy 480 Gen10 Plus")
        self.assertEqual(len(data["SmartStorage"]), 1)
        self.assertEqual(data["SmartStorage"][0]["PartNumber"], "P27544-001")
        self.assertEqual(data["SmartStorage"][0]["SerialNumber"], "PZCSU0ARHHP05V")
        self.assertEqual({disk["SerialNumber"] for disk in data["PhysicalDisks"]},
                         {"S711NE0W201627", "S711NE0W201630"})
        for disk in data["PhysicalDisks"]:
            self.assertEqual(disk["FirmwareVersion"], "HPK3")
            self.assertEqual(disk["InterfaceType"], "NVMe")
            self.assertAlmostEqual(disk["CapacityGB"], 480.103981056)
            self.assertEqual(disk["InterfaceSpeedMbps"], 16000)

    def test_real_battery_serial_and_single_battery_for_two_controllers(self):
        for raw, data in zip(self.raw, self.parsed):
            expected = raw["/redfish/v1/Chassis/1"]["Oem"]["Hpe"].get("SmartStorageBattery", [])
            rows = [row for row in inventory(data) if row[3] == "RAID батарея"]
            self.assertEqual([row[1] for row in rows], [b["SerialNumber"] for b in expected])
            self.assertEqual([row[2] for row in rows], [b["SparePartNumber"] for b in expected])
        self.assertEqual(self.parsed[3]["SmartStorage"][0]["SerialNumber"], "PEYJE0BRHA70K8")

    def test_adapter_part_numbers_are_taken_from_device_inventory(self):
        for raw, data in zip(self.raw, self.parsed):
            for adapter in data["NetworkAdapters"]:
                devices = [node for path, node in raw.items() if "/Devices/" in path and
                           node.get("SerialNumber", "").strip() == adapter["SerialNumber"]]
                self.assertEqual(adapter["PartNumber"], devices[0]["PartNumber"])

    def test_absence_is_filtered_for_every_component_category(self):
        for state in ("Absent", "NotPresent", "Not Installed", "Empty", "removed", " absent "):
            component = {"SerialNumber": "missing", "State": state, "LocationFormat": "PCISlot",
                         "BackupPowerSourceStatus": "Present"}
            data = {key: [component] for key in ("Processors", "NetworkAdapters", "SmartStorage",
                    "StorageBatteries", "Memory", "PhysicalDisks", "PowerSupplies", "StorageEnclosures",
                    "Fans", "USBDevices")}
            data.update(SDCard=state, TrustedModules=state, ProcessorCount=2, ProcessorModel="CPU")
            with self.subTest(state=state):
                self.assertEqual(inventory(data), [])
        self.assertEqual(inventory({"Memory": [{"DIMMStatus": "NotPresent"}]}), [])

    def test_installed_faulty_components_are_retained(self):
        data = {"Memory": [{"DIMMStatus": "MapOutError"}, {"DIMMStatus": "MapOutConfiguration"}],
                "PowerSupplies": [{"State": "UnavailableOffline", "Health": "Critical"}],
                "NetworkAdapters": [{"State": "Disabled", "Health": "Warning"}]}
        rows = inventory(data)
        self.assertEqual(len(rows), 4)
        self.assertEqual({row[12] for row in rows if row[3] == "Память"},
                         {"MapOutError", "MapOutConfiguration"})

    def test_absent_raw_storage_and_battery_do_not_reappear(self):
        raw = copy.deepcopy(self.raw[1])
        for path, data in raw.items():
            if path.endswith("/HostBusAdapters/0") or "/Controllers/0" in path or "/Drives/" in path:
                data["Status"] = {"State": "Absent"}
        self.assertEqual(get_array_controllers(raw), ([], []))
        raw = copy.deepcopy(self.raw[2])
        raw["/redfish/v1/Chassis/1"]["Oem"]["Hpe"]["SmartStorageBattery"][0]["Status"]["State"] = "Absent"
        data = dict(self.parsed[2], **get_chassis(raw))
        self.assertFalse(any(row[3] == "RAID батарея" for row in inventory(data)))

    def test_same_model_adapters_without_serials_in_different_slots_survive(self):
        base = "/redfish/v1/Systems/1/BaseNetworkAdapters/"
        raw = {base + str(i): {"Name": "NIC", "PartNumber": "123", "Location": f"Slot {i}"}
               for i in (1, 2)}
        raw[base + "3"] = {"Name": "NIC", "Status": {"State": "Absent"}}
        self.assertEqual(len(get_network_adapters(raw)), 2)

    def test_standard_storage_without_smartstorage_and_hba_disk_paths(self):
        raw = {path: node for path, node in self.raw[1].items() if "/SmartStorage/" not in path}
        controllers, disks = get_array_controllers(raw)
        self.assertEqual(len(controllers), 1)
        self.assertEqual(len(disks), 2)
        self.assertEqual(len(inventory({"SmartStorage": controllers})), 1)
        base = "/redfish/v1/Systems/1/SmartStorage/HostBusAdapters/12"
        controllers, disks = get_array_controllers({base: {"Model": "HBA"},
            base + "/DiskDrives/10": {"Model": "SSD", "SerialNumber": "disk"}})
        self.assertEqual(len(controllers), 1)
        self.assertEqual([d["SerialNumber"] for d in disks], ["disk"])


if __name__ == "__main__":
    unittest.main()
