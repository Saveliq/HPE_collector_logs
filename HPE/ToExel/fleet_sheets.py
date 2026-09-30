"""Server, firmware and Health summaries; the existing P/N sheet is untouched."""
from collections import Counter, defaultdict
from math import ceil

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from fleet_data import firmware_groups, health_issues, text


def _table(ws, name, headers, rows, start=1):
    for column, label in enumerate(headers, 1):
        cell = ws.cell(start, column, label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[start].height = 32
    for row_index, values in enumerate(rows, start + 1):
        height = 30
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value if value != "" else None)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            width = ws.column_dimensions[get_column_letter(column)].width
            height = max(height, 15 * ceil(len(str(value or "")) / max(width - 2, 1)))
        ws.row_dimensions[row_index].height = min(409, height)
    end = start + len(rows)
    ref = f"A{start}:{get_column_letter(len(headers))}{end}"
    if rows:
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name=None, showFirstColumn=False, showLastColumn=False,
                                             showRowStripes=False, showColumnStripes=False)
        ws.add_table(table)
    if start == 1 or name == "ServerSummary":
        ws.auto_filter.ref = ref
    return end


def _serials(servers, indices):
    return ", ".join(sorted(text(servers[i].get("SerialNumber")) or
                            f"Без S/N ({text(servers[i].get('SourceFile')) or i + 1})" for i in indices))


def _write_summary(wb, servers, issues):
    ws = wb.create_sheet("Сводка_серверов")
    for column, width in {"A": 36, "B": 34, "C": 28, "D": 21, "E": 20, "F": 85}.items():
        ws.column_dimensions[column].width = width
    metrics = [("Всего серверов", len(servers)),
               ("Моделей", len({text(s.get("Model")) for s in servers} - {""})),
               ("Health != OK", sum(bool(items) for items in issues)),
               ("Версий BIOS", len({text(s.get("BiosVersion")) for s in servers} - {""})),
               ("Версий iLO", len({text(s.get("iLOVersion")) for s in servers} - {""}))]
    for column, (label, value) in enumerate(metrics, 1):
        ws.cell(1, column, label).font = Font(bold=True)
        cell = ws.cell(2, column, value)
        cell.font = Font(bold=True)
        cell.number_format = "0"
    ws.row_dimensions[2].height = 30
    ws.merge_cells("A3:F3")
    ws["A3"] = "Health != OK: серверы с отклонениями Health/HealthRollup у сервера или компонента. Показатели — по всей выгрузке."
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 30
    ws["A5"] = "Модели и комбинации BIOS + iLO"
    ws["A5"].font = Font(bold=True)
    groups = defaultdict(list)
    models = Counter(text(s.get("Model")) for s in servers)
    for index, server in enumerate(servers):
        groups[tuple(text(server.get(field)) for field in ("Model", "BiosVersion", "iLOVersion"))].append(index)
    rows = [[model, bios, ilo, len(indices), len(indices) / models[model], _serials(servers, indices)]
            for (model, bios, ilo), indices in sorted(groups.items())]
    end = _table(ws, "ServerSummary", ["Модель сервера", "BIOS", "iLO", "Кол-во серверов", "% от модели", "S/N серверов"], rows, 6)
    for row in range(7, end + 1):
        ws.cell(row, 5).number_format = "0.0%"
    for label, field in (("BIOS", "BiosVersion"), ("iLO", "iLOVersion")):
        start = end + 4
        ws.cell(start - 1, 1, f"Независимая сводка {label}").font = Font(bold=True)
        independent = defaultdict(list)
        for index, server in enumerate(servers):
            independent[(text(server.get("Model")), text(server.get(field)))].append(index)
        rows = [[model, version, len(indices), models[model], len(indices) / models[model], _serials(servers, indices)]
                for (model, version), indices in sorted(independent.items())]
        end = _table(ws, "Summary" + label, ["Модель сервера", label, "Кол-во серверов", "Серверов в модели", "% от модели", "S/N серверов"], rows, start)
        for row in range(start + 1, end + 1):
            ws.cell(row, 5).number_format = "0.0%"
    ws.freeze_panes = "B7"


def write_fleet_sheets(workbook, servers):
    issues = [health_issues(server) for server in servers]
    _write_summary(workbook, servers, issues)
    ws = workbook.create_sheet("Микрокоды")
    for column, width in {"A": 34, "B": 27, "C": 52, "D": 25, "E": 14, "F": 16, "G": 85, "H": 45}.items():
        ws.column_dimensions[column].width = width
    rows = [[model, kind, identity, version, len(group["servers"]), len(group["components"]),
             _serials(servers, group["servers"]), role]
            for (model, kind, identity, version, role), group in sorted(firmware_groups(servers).items())]
    _table(ws, "ComponentFirmware", ["Модель сервера", "Тип компонента", "Модель / P/N", "Firmware",
                                    "Серверов", "Компонентов", "S/N серверов", "Назначение прошивки"], rows)
    ws.freeze_panes = "D2"

    ws = workbook.create_sheet("Health")
    for column, width in {"A": 22, "B": 34, "C": 32, "D": 24, "E": 24, "F": 25, "G": 18,
                          "H": 38, "I": 18, "J": 42, "K": 75, "L": 28}.items():
        ws.column_dimensions[column].width = width
    rows = []
    for server, server_issues in zip(servers, issues):
        for item in server_issues:
            rows.append([text(server.get("SerialNumber")), text(server.get("Model")), item["ComponentType"],
                         text(item.get("PartNumber") or item.get("SparePartNumber")), text(item.get("SerialNumber")),
                         item.get("State"), item.get("Health"), text(item.get("FirmwareVersion")),
                         item.get("HealthRollup"), text(item.get("Model") or item.get("Name")),
                         text(item.get("SourcePath")), text(server.get("SourceFile"))])
    _table(ws, "HealthIssues", ["S/N сервера", "Модель", "Компонент", "P/N", "S/N компонента", "State", "Health",
                               "Firmware", "HealthRollup", "Модель / имя компонента", "Путь Redfish", "Исходный файл"], rows)
    ws.freeze_panes = "C2"
    workbook.active = workbook.index(workbook["Сводка_серверов"])


def format_audit(ws):
    ws.freeze_panes = "E2"
    for column, width in {"A": 7, "B": 23, "C": 23, "D": 35, "E": 22, "F": 52, "G": 38,
                          "H": 38, "I": 25, "J": 18, "K": 18, "L": 22, "M": 30, "N": 30, "O": 22}.items():
        ws.column_dimensions[column].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    for row in ws.iter_rows(min_row=2):
        if row[0].value is None:
            continue
        ws.row_dimensions[row[0].row].height = 44
        for cell in row:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
