"""Excel presentation for diagnostic records; readings remain numeric."""
import json

from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter


GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
GOOD_STATUSES = {"ok", "enabled", "goodinuse", "up", "linkup", "operational"}
BAD_STATUSES = {"critical", "warning", "failed", "error", "degraded",
                "mapouterror", "mapoutconfiguration", "degradeda3dc"}
STATUS_COLUMNS = {"State", "Health", "HealthRollup", "DIMMStatus", "Дополнительный статус"}


def display_value(value):
    if value is None or (isinstance(value, str) and value.strip().lower() in {"", "none", "null", "n/a", "unknown", "нет данных"}):
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False) if value else None
    return value.strip() if isinstance(value, str) else value


def format_status_colors(workbook):
    """Color explicit status values only; missing and unknown states stay neutral."""
    for ws in workbook:
        columns = {cell.value: cell.column for cell in ws[1]}
        for row in ws.iter_rows(min_row=2):
            for label in STATUS_COLUMNS:
                if label not in columns:
                    continue
                cell = row[columns[label] - 1]
                status = str(cell.value or "").strip().lower()
                if status == "warning":
                    cell.fill = YELLOW_FILL
                elif status in BAD_STATUSES:
                    cell.fill = RED_FILL
                elif status in GOOD_STATUSES:
                    cell.fill = GREEN_FILL
        for label in STATUS_COLUMNS & columns.keys():
            if ws.max_row < 2:
                continue
            letter = get_column_letter(columns[label])
            for value, fill in (("Critical", RED_FILL), ("Warning", YELLOW_FILL)):
                ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}",
                                              CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill))
