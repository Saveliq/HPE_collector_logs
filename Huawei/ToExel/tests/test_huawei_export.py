import contextlib
import copy
import io
import json
import shutil
import sys
import unittest
import zipfile
from collections import Counter
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import ParseJsonToExel as export
from getAllJSONs import find_json_files
from getFromJSON.getChassis import get_chassis
from getFromJSON.getNetworkAdapters import get_network_adapters
from getFromJSON.getArrayControllers import get_array_controllers, get_storage_enclosures
from search import Resources

SYSTEM = "/redfish/v1/Systems/1"
CHASSIS = "/redfish/v1/Chassis/1"
STORAGE = SYSTEM + "/Storages/RAIDStorage0"
GROUPS = ("Processors", "Memory", "SmartStorage", "NetworkAdapters", "PhysicalDisks",
          "PowerSupplies", "Fans", "Boards", "StorageEnclosures", "StorageBatteries", "USBDevices")
EXPECTED_COUNTS = (2, 8, 1, 2, 6, 2, 4, 2, 1, 0, 0)


@contextlib.contextmanager
def output_folder():
    parent = Path(__file__).resolve().parent
    folder = parent / ("output_" + uuid4().hex)
    folder.mkdir()
    try:
        yield folder
    finally:
        # Only delete the unique test directory we just created.
        assert folder.resolve().parent == parent and folder.name.startswith("output_")
        shutil.rmtree(folder)


class HuaweiExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.file = ROOT / "examples/redfish_full.json"
        cls.raw = json.loads(cls.file.read_text(encoding="utf-8"))
        cls.data = export.parse_data(cls.raw)

    def test_sample_inventory_and_distinct_serials(self):
        self.assertEqual(tuple(len(self.data[k]) for k in GROUPS), EXPECTED_COUNTS)
        records = [node for k in GROUPS for node in self.data[k]]
        self.assertEqual(len(records), 28)
        serials = [node["SerialNumber"] for node in records if node.get("SerialNumber")]
        self.assertEqual(len(serials), 20)
        self.assertEqual(len(serials), len(set(serials)))
        self.assertNotIn("SDCard", self.data)
        self.assertNotIn("TrustedModules", self.data)

    def test_cpu_and_memory_fields_match_their_own_resources(self):
        for cpu in self.data["Processors"]:
            source = self.raw[cpu["SourcePath"]]
            for field in ("SerialNumber", "PartNumber"):
                self.assertEqual(cpu[field], source["Oem"]["Huawei"][field])
            self.assertEqual(cpu["Model"], source["Model"])
        for dimm in self.data["Memory"]:
            source = self.raw[dimm["SourcePath"]]
            for field in ("SerialNumber", "PartNumber"):
                self.assertEqual(dimm[field], source[field])
            for field, original in (("SizeMB", "CapacityMiB"), ("Frequency", "OperatingSpeedMhz"),
                                    ("Rank", "RankCount"), ("Technology", "MemoryDeviceType")):
                self.assertEqual(dimm[field], source[original])
            self.assertEqual(dimm["DIMMStatus"], source["Status"]["Health"])
        self.assertEqual(sum(n["SizeMB"] for n in self.data["Memory"]) / 1024,
                         self.raw[SYSTEM]["MemorySummary"]["TotalSystemMemoryGiB"])

    def test_controller_uses_associated_card_not_chassis(self):
        controller = self.data["SmartStorage"][0]
        self.assertEqual((controller["Model"], controller["PartNumber"], controller["SerialNumber"]),
                         ("SR150-M", "03024JNF", "024JNFCNJC010233"))
        source = self.raw[STORAGE]["StorageControllers"][0]
        self.assertEqual(controller["FirmwareVersion"], source["FirmwareVersion"])
        self.assertEqual(controller["CacheMemorySizeMiB"], 0)
        self.assertEqual(controller["CurrentOperatingMode"], "RAID")
        raw = copy.deepcopy(self.raw)
        raw[STORAGE]["StorageControllers"][0]["Links"] = {"Chassis": [{"@odata.id": CHASSIS}]}
        self.assertEqual(get_array_controllers(raw)[0][0]["PartNumber"], "03024JNF")

    def test_network_views_merge_and_missing_pcie_card_is_found(self):
        adapters = self.data["NetworkAdapters"]
        self.assertEqual([n["Name"] for n in adapters],
                         ["QLE2692-HUA-SP", "PRO/1000 PT Quad Port LP Server Adapter"])
        self.assertEqual(adapters[0]["PartNumber"], "QLE2692-HUA-SP")
        self.assertIsNone(adapters[1]["PartNumber"])
        self.assertTrue(all(n["SerialNumber"] is None and n["FirmwareVersion"] is None for n in adapters))
        mainboard = next(b for b in self.data["Boards"] if b["DeviceType"] == "MainBoard")
        self.assertEqual(mainboard["SerialNumber"], "024AFQ10K1000758")
        self.assertEqual(mainboard["PartNumber"], "03024AFQ")
        self.assertIn("X722", mainboard["Comment"])
        self.assertIn("2*10GE+2*GE", mainboard["Comment"])
        # An independently serialized LOM is a separate physical board.
        raw = copy.deepcopy(self.raw)
        raw[CHASSIS + "/Boards/mainboardLOM"]["SerialNumber"] = "SEPARATE-LOM"
        adapters = get_network_adapters(raw)
        self.assertEqual(len(adapters), 3)
        self.assertIn("SEPARATE-LOM", [n["SerialNumber"] for n in adapters])

    def test_backplane_has_own_identity_and_cpld(self):
        board = self.data["StorageEnclosures"][0]
        self.assertEqual(board["ComponentName"], "Дисковая корзина (backplane)")
        self.assertEqual(board["PartNumber"], "03022HXW")
        self.assertIsNone(board["SerialNumber"])
        self.assertEqual(board["FirmwareVersion"], "CPLD: 1.10")
        raw = copy.deepcopy(self.raw)
        raw[CHASSIS + "/Boards/chassisDiskBP1"]["Links"] = {
            "PCIeDevices": [{"@odata.id": CHASSIS + "/Boards/mainboardRAIDCard1"}]}
        self.assertEqual(get_storage_enclosures(raw)[0]["PartNumber"], "03022HXW")

    def test_disk_fields_and_psu_identity_are_not_inferred_from_parent(self):
        for disk in self.data["PhysicalDisks"]:
            source = self.raw[disk["SourcePath"]]
            self.assertEqual(disk["PartNumber"], source["Model"])
            self.assertEqual(disk["SerialNumber"], source["SerialNumber"])
            self.assertEqual(disk["FirmwareVersion"], source["Revision"])
            self.assertAlmostEqual(disk["CapacityGB"], source["CapacityBytes"] / 1e9)
            self.assertEqual(disk["InterfaceSpeedMbps"], source["NegotiatedSpeedGbs"] * 1000)
        supplies = self.raw[CHASSIS + "/Power"]["PowerSupplies"]
        for psu, source in zip(self.data["PowerSupplies"], supplies):
            for field in ("SerialNumber", "PartNumber", "FirmwareVersion", "PowerCapacityWatts"):
                self.assertEqual(psu[field], source[field])

    def test_fan_sensor_fallback_does_not_invent_inventory(self):
        self.assertNotIn(CHASSIS + "/Thermal", self.raw)
        self.assertEqual([n["Name"] for n in self.data["Fans"]], ["FAN1", "FAN2", "FAN3", "FAN4"])
        self.assertTrue(all(n["Reading"] == 2520 and n["Health"] == "ok" for n in self.data["Fans"]))
        raw = copy.deepcopy(self.raw)
        del raw[CHASSIS + "/ThresholdSensors"]
        self.assertEqual(get_chassis(raw)["Fans"], [])  # Slot limit and bit masks are insufficient.
        raw[CHASSIS + "/Thermal"] = {"Fans": [
            {"Name": "FAN1", "SerialNumber": "FAN-SN", "PartNumber": "FAN-PN", "Status": {"State": "Enabled"}},
            {"Name": "FAN2", "Status": {"State": "Absent"}}]}
        fans = get_chassis(raw)["Fans"]
        self.assertEqual(len(fans), 1)
        self.assertEqual(fans[0]["PartNumber"], "FAN-PN")

    def test_repeated_fragment_and_alias_responses_do_not_duplicate_components(self):
        raw = copy.deepcopy(self.raw)
        for path, node in list(raw.items()):
            raw[path + "#/duplicate"] = copy.deepcopy(node)
        path = CHASSIS + "/PCIeDevices/PCIeCard1"
        raw[CHASSIS + "/PCIeDevices/Alias1"] = copy.deepcopy(raw[path])
        raw[CHASSIS + "/Power"]["PowerSupplies"] *= 2
        data = export.parse_data(raw)
        self.assertEqual(tuple(len(data[k]) for k in GROUPS), EXPECTED_COUNTS)
        # Equal descriptions with no S/N are not proof that two slots are one card.
        path = CHASSIS + "/PCIeDevices/PCIeCard3"
        raw[path] = copy.deepcopy(self.raw[CHASSIS + "/PCIeDevices/PCIeCard2"])
        raw[path]["@odata.id"] = path
        raw[path]["Id"] = "PCIeCard3"
        self.assertEqual(len(get_network_adapters(raw)), 3)

    def test_multiple_views_of_one_unserialized_physical_adapter(self):
        raw = copy.deepcopy(self.raw)
        original = next(k for k in raw if k.startswith(CHASSIS + "/NetworkAdapters/") and "QLE" in k and "/NetworkPorts" not in k)
        alias = CHASSIS + "/NetworkAdapters/AnotherView"
        raw[alias] = copy.deepcopy(raw[original])
        raw[alias]["@odata.id"] = alias
        raw[alias]["Id"] = "AnotherView"
        self.assertEqual(len(get_network_adapters(raw)), 2)

    def test_missing_resources_and_error_bodies_do_not_create_hardware(self):
        for raw in ({}, [], {SYSTEM: {"Manufacturer": "HPE"}}, {SYSTEM: {"error": {}}}):
            self.assertIsNone(export.parse_data(raw))
        raw = {SYSTEM: {"Manufacturer": "Huawei", "Model": "Test"},
               CHASSIS + "/Drives/1": {"error": {"message": "unavailable"}}}
        data = export.parse_data(raw)
        self.assertTrue(all(data[k] == [] for k in GROUPS))
        self.assertIsNone(data["iBMCVersion"])

    def test_zero_false_and_failed_components_are_preserved(self):
        raw = copy.deepcopy(self.raw)
        raw[SYSTEM]["Oem"]["Huawei"]["PowerOnStrategy"] = False
        for path, node in raw.items():
            if node.get("@odata.type", "").endswith(".Processor"):
                node["Status"]["State"] = "Absent"
            if node.get("@odata.type", "").endswith(".Memory"):
                node["Status"] = {"State": "Disabled", "Health": "Critical"}
        raw[CHASSIS + "/Drives/HDDPlaneDisk0"]["CapacityBytes"] = 0
        raw[CHASSIS + "/Drives/HDDPlaneDisk0"]["NegotiatedSpeedGbs"] = 0
        data = export.parse_data(raw)
        self.assertIs(data["PowerAutoOn"], False)
        self.assertEqual(len(data["Memory"]), 8)
        self.assertTrue(all(d["DIMMStatus"] == "Critical" for d in data["Memory"]))
        self.assertEqual(data["PhysicalDisks"][0]["CapacityGB"], 0)
        self.assertEqual(data["PhysicalDisks"][0]["InterfaceSpeedMbps"], 0)
        ws = Workbook().active
        self.assertEqual(export.write_proc_info(ws, data, 1, 1)[0], 1)

    def test_current_firmware_and_power_policy_use_current_resources(self):
        data = self.data
        self.assertEqual(data["iBMCVersion"], "6.32")
        self.assertEqual(data["BiosVersion"], "8.20")
        self.assertEqual(data["PowerRegulatorMode"], "Custom")
        self.assertEqual(data["PowerAutoOn"], "TurnOn")
        raw = copy.deepcopy(self.raw)
        raw[SYSTEM + "/Bios/Settings"]["Attributes"]["CustomPowerPolicy"] = "WRONG-PENDING"
        self.assertEqual(export.parse_data(raw)["PowerRegulatorMode"], "Custom")

    def test_input_is_unchanged_and_resource_id_is_not_hardcoded(self):
        original = copy.deepcopy(self.raw)
        export.parse_data(self.raw)
        self.assertEqual(self.raw, original)
        text = json.dumps(original)
        for root in ("Systems", "Chassis", "Managers"):
            text = text.replace("/redfish/v1/" + root + "/1", "/redfish/v1/" + root + "/Blade9")
        data = export.parse_data(json.loads(text))
        self.assertEqual(tuple(len(data[k]) for k in GROUPS), EXPECTED_COUNTS)

    def test_excel_roundtrip_grouping_and_repeat_export(self):
        headers = ("№", "S/N", "P/N", "Наименование", "Spare P/N", "Description", "Quantity", "Комментарий",
                   "Версия прошивки", "State", "Health", "HealthRollup", "DIMMStatus", "Политика питания",
                   "Включение после потери питания")
        counts = {"Процессор": 2, "Память": 8, "RAID контроллер": 1, "Сетевой адаптер": 2, "Диск": 6,
                  "Блок питания": 2, "Вентилятор": 4, "Системная плата": 1, "PCIe райзер": 1,
                  "Дисковая корзина (backplane)": 1}
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            for _ in range(2):
                export._parseJSONToExel([self.file, self.file], folder)
                wb = load_workbook(folder / export.OUTPUT_PATH)
                try:
                    self.assertEqual(wb.sheetnames, ["Аудит", "Группировка_PN"])
                    ws = wb["Аудит"]
                    self.assertEqual(next(ws.values), headers)
                    self.assertEqual(ws["B2"].value, "2102311XBK10K1001017")
                    self.assertEqual(ws["C2"].value, "02311XBK")
                    self.assertEqual(ws["I2"].value, "BIOS: 8.20; iBMC: 6.32")
                    rows = [row for row in ws.values if row[6] == 1]
                    self.assertEqual(Counter(row[3] for row in rows), counts)
                    self.assertEqual(len([r for r in rows if r[1] == "024AFQ10K1000758"]), 1)
                    self.assertEqual(len([r for r in rows if r[1] == "024JNFCNJC010233"]), 1)
                    self.assertEqual(len([r for r in ws.values if isinstance(r[0], int)]), 1)
                    summary = {r[0]: r[2] for r in wb["Группировка_PN"].values if isinstance(r[2], int)}
                    self.assertEqual(sum(summary.values()), 28)
                    self.assertEqual(summary["36ASF4G72PZ-2G6D1"], 8)
                    self.assertEqual(summary["AL14SEB060N"], 4)
                    self.assertEqual(summary["03024JNF"], 1)
                    for row in ws.iter_rows(min_row=2):
                        if row[3].value == "Память":
                            self.assertEqual(row[12].value, "OK")
                            self.assertEqual(row[12].fill.fgColor.rgb, "00C6EFCE")
                    self.assertEqual(ws.column_dimensions["I"].width, 30)
                    self.assertEqual(ws.row_dimensions[1].height, 42)
                finally:
                    wb.close()

    def test_zip_files_with_identical_member_names_preserve_both_servers(self):
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            for index in (1, 2):
                raw = copy.deepcopy(self.raw)
                raw[SYSTEM]["SerialNumber"] = "SERVER-" + str(index)
                with zipfile.ZipFile(folder / (str(index) + ".zip"), "w") as archive:
                    archive.writestr("redfish_full.json", json.dumps(raw))
            tmp = folder / "tmp"
            files = find_json_files(folder, tmp)
            self.assertEqual(len(files), 2)
            self.assertEqual(len(set(files)), 2)
            export._parseJSONToExel(files, folder)
            wb = load_workbook(folder / export.OUTPUT_PATH)
            try:
                self.assertEqual({r[1] for r in wb["Аудит"].values if isinstance(r[0], int)}, {"SERVER-1", "SERVER-2"})
            finally:
                wb.close()
            self.assertEqual(len(find_json_files(folder, tmp)), 2)

    def test_invalid_json_and_missing_file_are_skipped(self):
        with output_folder() as folder, contextlib.redirect_stdout(io.StringIO()):
            path = folder / "invalid.json"
            path.write_text("not JSON", encoding="utf-8")
            self.assertIsNone(export.print_parser(path))
            self.assertIsNone(export.print_parser(folder / "missing.json"))
            path.write_text(json.dumps(self.raw), encoding="utf-8-sig")
            self.assertEqual(export.print_parser(path)["Model"], "2288H V5")

    def test_all_project_imports_are_local(self):
        import main
        self.assertIs(main._parseJSONToExel, export._parseJSONToExel)
        modules = ["main", "ParseJsonToExel", "search", "diagnostic_sheets", "getAllJSONs"]
        modules += ["getFromJSON." + p.stem for p in (ROOT / "getFromJSON").glob("*.py")]
        for name in modules:
            self.assertTrue(Path(sys.modules[name].__file__).resolve().is_relative_to(ROOT), name)


if __name__ == "__main__":
    unittest.main()
